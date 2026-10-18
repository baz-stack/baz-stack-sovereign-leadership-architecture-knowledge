#!/usr/bin/env python3
"""Text-completeness check: every paragraph and table cell of a DOCX must appear in its Markdown conversion.
Usage: python3 -I tools/check_conversion.py <file.docx> <file.md>"""
import re
import sys
import docx


def norm(s):
    s = s.replace("\\", "").replace("*", "").replace("_", " ")
    s = s.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    return re.sub(r"[^a-z0-9$%]+", "", s.lower())


def units(path):
    d = docx.Document(path)
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                out += [p.text for p in c.paragraphs if p.text.strip()]
    return out


def main():
    src, md = sys.argv[1], open(sys.argv[2], encoding="utf-8").read()
    hay = norm(md)
    us = units(src)
    miss = [u for u in us if norm(u) and norm(u) not in hay]
    print(f"{src.split('/')[-1]}: {len(us)} text units, {len(miss)} missing from Markdown")
    for m in miss[:10]:
        print("  MISSING:", m[:100])
    return 1 if miss else 0


if __name__ == "__main__":
    sys.exit(main())
